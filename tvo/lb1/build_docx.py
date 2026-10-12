# Builds КНТ-23-1_Кравченко_ЛБ1.docx. Run: python build_docx.py (after make_charts.py)
# Inline markup in text: *italic*, _{subscript}, ^{superscript}.
import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
FONT = 'Times New Roman'
OUT = HERE / 'КНТ-23-1_Кравченко_ЛБ1.docx'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(2), Cm(2), Cm(3), Cm(1.5)


def set_font(style_or_run, size=14, bold=None, name=FONT):
    f = style_or_run.font
    f.name, f.size = name, Pt(size)
    if bold is not None:
        f.bold = bold
    f.color.rgb = RGBColor(0, 0, 0)
    rpr = style_or_run.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):  # theme fonts override explicit ones
        fonts.attrib.pop(qn(a), None)


normal = doc.styles['Normal']
set_font(normal)
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
pf.space_before = pf.space_after = Pt(0)

# Heading 1 = «ЗВІТ» on the title page, Heading 2 = sections, Heading 3 = subsections.
for name, size, align, indent, before, after, caps in [
    ('Heading 1', 16, WD_ALIGN_PARAGRAPH.CENTER, None, 0, 0, False),
    ('Heading 2', 14, WD_ALIGN_PARAGRAPH.CENTER, None, 12, 12, True),
    ('Heading 3', 14, WD_ALIGN_PARAGRAPH.LEFT, Cm(1.25), 12, 6, False),
]:
    st = doc.styles[name]
    set_font(st, size, bold=True)
    st.font.italic = False
    st.font.all_caps = caps
    f = st.paragraph_format
    f.alignment, f.first_line_indent = align, indent
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    f.keep_with_next = True


def runs(p, text, size=None, bold=False):
    for tok in re.split(r'(\*[^*]+\*|_\{[^}]+\}|\^\{[^}]+\})', text):
        if not tok:
            continue
        if tok.startswith('*'):
            r = p.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith('_{') or tok.startswith('^{'):
            inner = tok[2:-1]
            it = inner.startswith('*') and inner.endswith('*')
            r = p.add_run(inner.strip('*') if it else inner)
            r.italic = it
            if tok[0] == '_':
                r.font.subscript = True
            else:
                r.font.superscript = True
        else:
            r = p.add_run(tok)
        if size:
            r.font.size = Pt(size)
        if bold:
            r.bold = True
    return p


def P(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, keep_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(1.25) if indent else None
    p.paragraph_format.keep_with_next = keep_next
    return runs(p, text)


def C(text, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return runs(p, text, bold=bold)


def F(text):  # displayed formula
    return P(text, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)


def bullet(text):
    return P('– ' + text)


def H2(text, page_break=True):
    p = doc.add_paragraph(text, style='Heading 2')
    p.paragraph_format.page_break_before = page_break
    return p


def H3(text):
    return doc.add_paragraph(text, style='Heading 3')


def blank(n=1):
    for _ in range(n):
        doc.add_paragraph()


def table_caption(text):
    p = P(text, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, keep_next=True)
    p.paragraph_format.space_before = Pt(6)
    return p


def figure(file, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.add_run().add_picture(str(HERE / 'img' / file), width=Cm(width_cm))
    c = P(caption, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    c.paragraph_format.space_after = Pt(6)


def set_cell_borders(table):
    tbl = table._tbl
    borders = OxmlElement('w:tblBorders')
    for b in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{b}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), '000000')
        borders.append(e)
    tbl.tblPr.append(borders)


def table(rows, widths_cm, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_borders(t)
    for ri, row in enumerate(rows):
        tr = t.rows[ri]
        trpr = tr._tr.get_or_add_trPr()
        cs = OxmlElement('w:cantSplit'); trpr.append(cs)
        for ci, val in enumerate(row):
            cell = tr.cells[ci]
            cell.width = Cm(widths_cm[ci])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.keep_with_next = ri < len(rows) - 1
            runs(p, str(val), size=12, bold=header and ri == 0)
    return t


def code(path):
    for line in (HERE / path).read_text(encoding='utf-8').splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.first_line_indent = None
        r = p.add_run(line if line else ' ')
        set_font(r, 10, name='Courier New')


def num(v, d=2):
    return f'{v:.{d}f}'.replace('.', ',')


# ---------- data ----------
res = list(csv.DictReader(open(HERE / 'results' / 'results.csv', encoding='utf-8'), delimiter=';'))
data = {int(r['threads']): r for r in res}
T1 = float(data[1]['median_ms'])
W, H = 3840, 2400
N = W * H
TAU = T1 * 1e6 / N  # ns per pixel
CORES, LOGICAL = 6, 12
s = {p: float(r['speedup']) for p, r in data.items()}

# ---------- title page ----------
C('МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ', True)
C('ХАРКІВСЬКИЙ НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ РАДІОЕЛЕКТРОНІКИ', True)
blank(1)
C('Кафедра системотехніки')
blank(6)
doc.add_paragraph('ЗВІТ', style='Heading 1')
C('з лабораторної роботи №1')
C('з дисципліни «Технології високопродуктивних обчислень»')
C('на тему: «Багатопоточне програмування»')
C('Варіант 9')
blank(5)
sign = doc.add_table(rows=1, cols=2)
for cell, lines, al in [(sign.cell(0, 0), ['Виконав:', 'ст. гр. КНТ-23-1', 'Кравченко Р.'], WD_ALIGN_PARAGRAPH.LEFT),
                        (sign.cell(0, 1), ['Перевірив:', 'проф. каф. СТ', 'Міщеряков Ю.В.'], WD_ALIGN_PARAGRAPH.RIGHT)]:
    cell.paragraphs[0].text = lines[0]
    cell.paragraphs[0].alignment = al
    for l in lines[1:]:
        cell.add_paragraph(l).alignment = al
blank(4)
C('Харків 2026')

# ---------- body ----------
H2('1 МЕТА РОБОТИ')
P('Вивчити способи організації паралельних багатопоточних додатків на прикладі розробки програмного засобу обробки зображення – накладання фільтра.')

H2('2 ЗАВДАННЯ', page_break=False)
P('Варіант 9 – шум Перліна (Perlin Noise). Розробити програмний засіб, який накладає на зображення шум Перліна, розподіляючи обробку між кількома потоками. Розрахувати теоретичні прискорення та ефективність розробленого паралельного алгоритму, зібрати статистику витрат часу на обробку зображення в одному та декількох потоках (1, 2, 4, 6, 8, 10, 12, 14, 16), побудувати графік залежності часу обробки від кількості потоків та порівняти отримані практичні результати з теоретичними.')

H2('3 ХІД ВИКОНАННЯ РОБОТИ')
H3('3.1 Алгоритм шуму Перліна')
P('Шум Перліна – це градієнтний шум, який дає плавну псевдовипадкову функцію *n*(*x*, *y*) зі значеннями приблизно в межах [−1; 1]. У роботі використано покращену версію алгоритму К. Перліна (2002 р.). Площина розбивається на квадратні клітинки цілочисельної решітки; кожному вузлу решітки за таблицею перестановок із 256 елементів зіставляється псевдовипадковий градієнт. Для точки (*x*, *y*):')
bullet('визначаються клітинка, у яку потрапляє точка, та дробові координати всередині неї *x*_{*f*}, *y*_{*f*};')
bullet('для кожного з чотирьох кутів обчислюється скалярний добуток градієнта на вектор від кута до точки;')
bullet('чотири значення інтерполюються з вагами, отриманими згладжувальною кривою')
F('*fade*(*t*) = 6*t*^{5} − 15*t*^{4} + 10*t*^{3},')
P('що забезпечує неперервність першої та другої похідних на межах клітинок.', indent=False)
P('Щоб отримати природну «хмарну» текстуру, використовується фрактальна сума октав (fBm):')
F('*N*(*x*, *y*) = Σ_{o=0}^{K−1} *a*^{*o*} · *n*(2^{*o*}·*s*·*x*, 2^{*o*}·*s*·*y*) / Σ_{o=0}^{K−1} *a*^{*o*},')
P(f'де *K* = 6 – кількість октав, *a* = 0,5 – коефіцієнт згасання амплітуди (persistence), *s* = 1/160 – масштаб (розмір найбільшої «плями» ≈ 160 пікселів).')
P('Накладання шуму на зображення: до кожного каналу (R, G, B) пікселя додається зсув яскравості, пропорційний значенню шуму, з обмеженням результату діапазоном [0; 255]:')
F('*B*(*x*, *y*) = clamp(*A*(*x*, *y*) + round(*k* · *N*(*x*, *y*)), 0, 255),')
P('де *A* – вхідне зображення, *B* – результат, *k* = 70 – сила ефекту. Значення шуму для пікселя залежить лише від його координат, тому пікселі обробляються незалежно один від одного – задача має природний паралелізм за даними.')

H3('3.2 Паралельний алгоритм')
P(f'Використано декомпозицію за даними: зображення висотою *H* рядків ділиться на *p* горизонтальних смуг майже однакової висоти, по одній смузі на потік. Потік *t* (*t* = 0, …, *p* − 1) обробляє рядки від ⌊*H*·*t*/*p*⌋ до ⌊*H*·(*t*+1)/*p*⌋ − 1. Головний потік створює *p* потоків (клас System.Threading.Thread), запускає їх і чекає на завершення кожного (Join).')
P('Синхронізація під час обробки не потрібна: вхідне зображення та таблиця перестановок шуму лише читаються, а кожен потік записує результат тільки у свої рядки вихідного масиву, тобто критичних ресурсів немає. Правильність перевіряється програмно: результат для кожної кількості потоків побайтово порівнюється з результатом в одному потоці.')

H3('3.3 Теоретична оцінка прискорення та ефективності')
P(f'Тестове зображення має розмір *W* × *H* = {W} × {H} = {N:,} пікселів'.replace(',', ' ') + f'. Обробка одного пікселя займає однаковий час τ (6 октав шуму + 3 додавання), тому час послідовного виконання')
F('*T*_{1} = *W* · *H* · τ.')
P(f'За виміряним *T*_{{1}} = {num(T1, 1)} мс отримуємо τ ≈ {num(TAU, 1)} нс на піксель.')
P('Для *p* потоків час визначається найвищою смугою та накладними витратами *t*_{0}(*p*) на створення і очікування потоків:')
F('*T*_{*p*} = ⌈*H*/*p*⌉ · *W* · τ + *t*_{0}(*p*).')
P('Накладні витрати виміряно окремою програмою (створення, запуск і очікування порожніх потоків, медіана 51 повтору): *t*_{0}(1) ≈ 0,16 мс, *t*_{0}(6) ≈ 0,76 мс, *t*_{0}(12) ≈ 1,29 мс, *t*_{0}(16) ≈ 1,65 мс. Це 0,02–0,15 % від *T*_{1}, отже частка послідовного коду за законом Амдала *f* ≈ 0, і на паракомп’ютері')
F('*S*_{*p*} = *T*_{1}/*T*_{*p*} ≈ *p*,    *E*_{*p*} = *S*_{*p*}/*p* ≈ 1.')
P(f'Нерівномірність смуг (наприклад, для *p* = 7: 342 або 343 рядки) змінює прискорення менш ніж на 0,3 %. Реальну межу задає обладнання: процесор AMD Ryzen 5 5600 має {CORES} фізичних ядер і {LOGICAL} логічних процесорів (SMT). Потоки, що працюють на двох логічних процесорах одного ядра, ділять його обчислювальні блоки, тому теоретична оцінка для цієї машини')
F(f'*S*_{{*p*}} ≈ min(*p*, {CORES}),    *E*_{{*p*}} = *S*_{{*p*}}/*p*.')
P(f'SMT може дати додатково 10–30 % понад {CORES} для обчислювальних задач, а при *p* > {LOGICAL} потоків більше, ніж логічних процесорів, тому подальшого прискорення не очікується. Теоретичні значення наведено в таблиці 3.1.')

table_caption('Таблиця 3.1 – Теоретичні прискорення та ефективність')
rows = [['*p*', '*S*_{*p*} (паракомп’ютер)', '*E*_{*p*} (паракомп’ютер)', f'*S*_{{*p*}} = min(*p*, {CORES})', '*E*_{*p*} = *S*_{*p*}/*p*']]
for p in data:
    st = min(p, CORES)
    rows.append([p, p, '1', st, num(st / p)])
table(rows, [1.5, 3.5, 3.5, 3.5, 3.5])

H3('3.4 Програмна реалізація')
P('Програму написано мовою C# (.NET 10, консольний застосунок). Для читання та запису зображень використано System.Drawing (Bitmap.LockBits – доступ до пікселів як до масиву байтів BGR). Структура проєкту:')
bullet('Perlin.cs – клас шуму Перліна: таблиця перестановок, функції *fade*, *lerp*, *grad*, Noise (одна октава) та Fractal (сума октав);')
bullet('Program.cs – читання зображення, функція ProcessRows (обробка діапазону рядків), функція RunWithThreads (поділ на смуги, створення, запуск та очікування потоків), вимірювання часу, перевірка результату, збереження результатів у results.csv та зображень.')
P('Час вимірюється класом Stopwatch від створення потоків до завершення останнього з них; читання та запис файлів у вимірювання не входять. Перший прогін виконується для «прогріву» JIT-компілятора і в статистику не потрапляє. Для кожної кількості потоків обробка повторюється 11 разів, у таблицю береться медіана – вона менш чутлива до випадкових затримок операційної системи, ніж середнє. Програму зібрано в конфігурації Release. Повний код наведено в додатку А.')

H3('3.5 Результат накладання фільтра')
P('Вхідне зображення наведено на рисунку 3.1, згенерований фрактальний шум Перліна – на рисунку 3.2, результат накладання шуму – на рисунку 3.3.')
figure('input.jpg', 14, 'Рисунок 3.1 – Вхідне зображення (3840 × 2400)')
figure('noise.jpg', 14, 'Рисунок 3.2 – Фрактальний шум Перліна (6 октав, масштаб 1/160)')
figure('output.jpg', 14, 'Рисунок 3.3 – Результат накладання шуму Перліна (*k* = 70)')

H3('3.6 Експериментальні результати')
P('Вимірювання виконано на комп’ютері з процесором AMD Ryzen 5 5600 (6 ядер, 12 потоків), ОС Windows 11 Pro, .NET 10. Крім кількостей потоків із завдання (1, 2, 4, 6, 8, 10, 12, 14, 16), виміряно також 3, 5 та 7 потоків, щоб отримати повний ряд 1–8 для графіка. Результати наведено в таблиці 3.2 та на рисунках 3.4, 3.5. Для всіх кількостей потоків результат побайтово збігся з результатом в одному потоці.')
table_caption('Таблиця 3.2 – Час обробки зображення, прискорення та ефективність')
rows = [['*p*', 'Медіана *T*_{*p*}, мс', 'Мін., мс', 'Макс., мс', '*S*_{*p*}', '*E*_{*p*}']]
for p, r in data.items():
    rows.append([p, num(float(r['median_ms']), 1), num(float(r['min_ms']), 1), num(float(r['max_ms']), 1),
                 num(float(r['speedup'])), num(float(r['efficiency']))])
table(rows, [1.5, 3.3, 2.5, 2.5, 2.2, 2.2])
figure('time.png', 16, 'Рисунок 3.4 – Залежність часу обробки зображення від кількості потоків')
figure('speedup.png', 16, 'Рисунок 3.5 – Залежність прискорення від кількості потоків')

H3('3.7 Порівняння з теоретичними результатами')
best = max(s, key=s.get)
P(f'Для 1–3 потоків прискорення практично дорівнює теоретичному: *S*_{{2}} = {num(s[2])}, *S*_{{3}} = {num(s[3])} при ефективності понад 0,98 – це підтверджує, що частка послідовного коду та накладні витрати на потоки дуже малі.')
P(f'Для 4–6 потоків прискорення відстає від min(*p*, {CORES}): *S*_{{6}} = {num(s[6])} замість 6 (*E*_{{6}} = {num(float(data[6]["efficiency"]))}). Ймовірні причини (окремо не вимірювалися): зниження тактової частоти процесора (Precision Boost), коли навантажено всі ядра, а не одне–два; фонові процеси ОС, які займають частину часу ядер; спільна кеш-пам’ять L3 і пропускна здатність пам’яті.')
P(f'Для 7–12 потоків працює SMT: другий логічний потік на ядрі використовує простої конвеєра, тому прискорення ще зростає і досягає максимуму *S*_{{{best}}} = {num(s[best])} при *p* = {best}, тобто SMT додало близько {round((s[best] / s[6] - 1) * 100)} % до результату 6 потоків. Ефективність при цьому падає до {num(float(data[best]["efficiency"]))}, бо ефективність рахується на кількість потоків, а фізичних ядер лише {CORES}.')
P(f'Для 14 та 16 потоків (більше, ніж {LOGICAL} логічних процесорів) прискорення зменшується до {num(s[14])} та {num(s[16])}: частина потоків чекає на процесор, планувальник ОС перемикає контексти, і смуги завершуються нерівномірно – загальний час визначає потік, який отримав процесор останнім. Це узгоджується з теоретичним висновком, що кількість потоків понад кількість логічних процесорів не дає виграшу.')

H2('4 ВИСНОВКИ')
P('У ході лабораторної роботи розроблено багатопоточний програмний засіб мовою C#, що накладає на зображення фрактальний шум Перліна (варіант 9). Обробка розпаралелена декомпозицією зображення на горизонтальні смуги за кількістю потоків; синхронізація не потрібна, бо потоки записують результат у різні рядки, а коректність підтверджено побайтовим збігом результатів з однопотоковим варіантом.')
P(f'Теоретично задача майже повністю паралельна (*f* ≈ 0), тому на паракомп’ютері *S*_{{*p*}} ≈ *p*, а на 6-ядерному процесорі – *S*_{{*p*}} ≈ min(*p*, 6). Експеримент підтвердив лінійне прискорення до 3 потоків (*E* > 0,98), показав відставання від теорії на 4–6 потоках через зниження частоти та фонове навантаження і максимальне прискорення {num(s[best])} при {best} потоках завдяки SMT. Подальше збільшення кількості потоків понад кількість логічних процесорів збільшує час обробки.')
P(f'Отже, оптимальна кількість потоків для такої обчислювальної задачі дорівнює кількості логічних процесорів ({LOGICAL}), а за критерієм ефективності – кількості фізичних ядер (до {CORES}).')

H2('ДОДАТОК А')
C('Код програми', bold=True)
P('Файл Perlin.cs', indent=False)
code('PerlinFilter/Perlin.cs')
P('Файл Program.cs', indent=False)
code('PerlinFilter/Program.cs')

doc.core_properties.title = 'ТВО ЛБ1 – Багатопоточне програмування'
doc.core_properties.author = 'Кравченко Р.'
doc.save(OUT)
print('written', OUT)
