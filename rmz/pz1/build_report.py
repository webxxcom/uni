# Builds КНТ-23-1_Кравченко_ПЗ1.docx (РмЗ). Run: python build_report.py
# Inline markup in text: *italic*, _{subscript}, ^{superscript}.
import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
FONT = 'Times New Roman'
OUT = HERE / 'КНТ-23-1_Кравченко_ПЗ1.docx'
PROJ = HERE / 'PhotoGallery'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(2), Cm(2), Cm(3), Cm(1.5)


def set_font(style_or_run, size=14, bold=None, name=FONT):
    f = style_or_run.font
    f.name, f.size = name, Pt(size)
    if bold is not None:
        f.bold = bold
    f.color.rgb = RGBColor(0, 0, 0)
    rpr = style_or_run.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):  # theme fonts override explicit ones
        fonts.attrib.pop(qn(a), None)


normal = doc.styles['Normal']
set_font(normal)
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
pf.space_before = pf.space_after = Pt(0)

# Heading 1 = «ЗВІТ» on the title page, Heading 2 = sections, Heading 3 = subsections.
for name, size, align, indent, before, after, caps in [
    ('Heading 1', 16, WD_ALIGN_PARAGRAPH.CENTER, None, 0, 0, False),
    ('Heading 2', 14, WD_ALIGN_PARAGRAPH.CENTER, None, 12, 12, True),
    ('Heading 3', 14, WD_ALIGN_PARAGRAPH.LEFT, Cm(1.25), 12, 6, False),
]:
    st = doc.styles[name]
    set_font(st, size, bold=True)
    st.font.italic = False
    st.font.all_caps = caps
    f = st.paragraph_format
    f.alignment, f.first_line_indent = align, indent
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    f.keep_with_next = True


def runs(p, text, size=None, bold=False):
    for tok in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|_\{[^}]+\}|\^\{[^}]+\})', text):
        if not tok:
            continue
        if tok.startswith('**'):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('*'):
            r = p.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith('_{') or tok.startswith('^{'):
            inner = tok[2:-1]
            it = inner.startswith('*') and inner.endswith('*')
            r = p.add_run(inner.strip('*') if it else inner)
            r.italic = it
            if tok[0] == '_':
                r.font.subscript = True
            else:
                r.font.superscript = True
        else:
            r = p.add_run(tok)
        if size:
            r.font.size = Pt(size)
        if bold:
            r.bold = True
    return p


def P(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, keep_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(1.25) if indent else None
    p.paragraph_format.keep_with_next = keep_next
    return runs(p, text)


def C(text, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return runs(p, text, bold=bold)


def F(text):  # displayed formula
    return P(text, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)


def bullet(text):
    return P('- ' + text)


def H2(text, page_break=True):
    p = doc.add_paragraph(text, style='Heading 2')
    p.paragraph_format.page_break_before = page_break
    return p


def H3(text):
    return doc.add_paragraph(text, style='Heading 3')


def blank(n=1):
    for _ in range(n):
        doc.add_paragraph()


def table_caption(text):
    p = P(text, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, keep_next=True)
    p.paragraph_format.space_before = Pt(6)
    return p


def figure(file, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.add_run().add_picture(str(HERE / 'img' / file), width=Cm(width_cm))
    c = P(caption, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    c.paragraph_format.space_after = Pt(6)


def set_cell_borders(table):
    tbl = table._tbl
    borders = OxmlElement('w:tblBorders')
    for b in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{b}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), '000000')
        borders.append(e)
    tbl.tblPr.append(borders)


def table(rows, widths_cm, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_borders(t)
    for ri, row in enumerate(rows):
        tr = t.rows[ri]
        trpr = tr._tr.get_or_add_trPr()
        cs = OxmlElement('w:cantSplit'); trpr.append(cs)
        for ci, val in enumerate(row):
            cell = tr.cells[ci]
            cell.width = Cm(widths_cm[ci])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.keep_with_next = ri < len(rows) - 1
            runs(p, str(val), size=12, bold=header and ri == 0)
    return t


def code(path, size=9):
    for line in (PROJ / path).read_text(encoding='utf-8').splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.first_line_indent = None
        r = p.add_run(line if line else ' ')
        set_font(r, size, name='Courier New')


def num(v, d=2):
    return f'{v:.{d}f}'.replace('.', ',')



exec((HERE / 'report_body.py').read_text(encoding='utf-8'))
